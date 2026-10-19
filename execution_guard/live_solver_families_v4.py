"""Fresh-goal EXECUTION families, wave 4 (heartbeat).  Same contract as waves 2 & 3: fresh inputs,
one plain sentence, an oracle that judges the real artifact with the standard library or typed readers.

  DOCX_CREATE              docx.document.create.python_docx
  XLSX_CREATE              xlsx.table.create_from_json_records
  XLSX_VERIFY              xlsx.table.verify_against_json_records
  SOURCE_RANKING_BM25      source.objective_relevance.rank.bm25
  DECISION_SYNTHESIS       decision.synthesis.typed.stdlib
  PYTHON_SOURCE_AUDIT      python.source.audit.ast
"""
from __future__ import annotations

import json
import random
from pathlib import Path
from typing import Any, Callable

from execution_guard.live_solver_end_to_end_measurement_v1 import Case, ROOT, _rel, _word
from execution_guard.live_solver_families_v2 import _passed


# ----------------------------------------------------------------------------- docx create

def gen_docx_create(rng: random.Random, work: Path, i: int) -> Case:
    work.mkdir(parents=True, exist_ok=True)
    title = "Report " + _word(rng, 6)
    para = "Paragraph " + _word(rng, 10)
    out = work / f"doc_{i}_{_word(rng, 4)}.docx"
    goal = rng.choice([
        f"Generate a docx document to {_rel(out)} with title {title}, paragraph {para}, table with headers Name, Status containing Alpha|Active and Beta|Done",
        f"Create a docx document at {_rel(out)} with title {title}, paragraph {para}, table with headers Name, Status containing Alpha|Active and Beta|Done",
        f"Generate a docx document to {_rel(out)} with title {title}, paragraph {para}, table with columns Name, Status containing Alpha|Active and Beta|Done",
    ])

    def oracle(result: dict[str, Any]) -> tuple[bool, str]:
        if not _passed(result):
            return False, "runtime did not pass: " + str(result.get("status"))
        if not out.is_file():
            return False, "docx output not written"
        raw = out.read_bytes()
        if not raw.startswith(b"PK\x03\x04"):
            return False, "docx file lacks ZIP magic header"
        import docx
        doc = docx.Document(out)
        headings = [p.text for p in doc.paragraphs if p.text == title]
        if not headings:
            return False, f"docx title {title!r} not found in document"
        return True, ""
    return Case("DOCX_CREATE", goal, oracle, True, [out])


# ----------------------------------------------------------------------------- xlsx create & verify

def gen_xlsx_create(rng: random.Random, work: Path, i: int) -> Case:
    work.mkdir(parents=True, exist_ok=True)
    rows = [{"name": _word(rng, 5), "score": rng.randint(10, 99)} for _ in range(rng.randint(2, 6))]
    src = work / f"records_{i}_{_word(rng, 4)}.json"
    src.write_text(json.dumps({"items": rows}), encoding="utf-8")
    out = work / f"table_{i}_{_word(rng, 4)}.xlsx"
    goal = rng.choice([
        f"Create an xlsx table from {_rel(src)} and save it to {_rel(out)} with columns name, score.",
        f"Create an xlsx table from {_rel(src)} and write it to {_rel(out)} with columns name, score.",
        f"Build an xlsx table from {_rel(src)} and save it to {_rel(out)} with columns name, score.",
    ])

    def oracle(result: dict[str, Any]) -> tuple[bool, str]:
        if not _passed(result):
            return False, "runtime did not pass: " + str(result.get("status"))
        if not out.is_file():
            return False, "xlsx output not written"
        raw = out.read_bytes()
        if not raw.startswith(b"PK\x03\x04"):
            return False, "xlsx file lacks ZIP magic header"
        import openpyxl
        wb = openpyxl.load_workbook(out)
        sheet = wb.active
        names = [sheet.cell(r, 1).value for r in range(2, sheet.max_row + 1)]
        expected_names = [r["name"] for r in rows]
        if names != expected_names:
            return False, f"xlsx rows mismatch: {names} != {expected_names}"
        return True, ""
    return Case("XLSX_CREATE", goal, oracle, True, [out])


def gen_xlsx_verify(rng: random.Random, work: Path, i: int) -> Case:
    work.mkdir(parents=True, exist_ok=True)
    rows = [{"name": _word(rng, 5), "score": rng.randint(10, 99)} for _ in range(rng.randint(2, 6))]
    src = work / f"vrecords_{i}_{_word(rng, 4)}.json"
    src.write_text(json.dumps({"items": rows}), encoding="utf-8")
    out = work / f"vtable_{i}_{_word(rng, 4)}.xlsx"
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "data"
    ws.cell(1, 1, "name")
    ws.cell(1, 2, "score")
    tamper = (i % 3 == 2)
    for r_idx, r in enumerate(rows, 2):
        ws.cell(r_idx, 1, r["name"] + ("_tampered" if tamper and r_idx == 2 else ""))
        ws.cell(r_idx, 2, r["score"])
    wb.save(out)
    goal = rng.choice([
        f"Verify the xlsx table {_rel(out)} against the json records in {_rel(src)} with columns name, score.",
        f"Check the xlsx table {_rel(out)} against the json records in {_rel(src)} with columns name, score.",
        f"Verify the xlsx table {_rel(out)} against the json records in {_rel(src)}, with columns name, score.",
    ])

    def oracle(result: dict[str, Any]) -> tuple[bool, str]:
        if tamper:
            return (not _passed(result)), "tampered table must be refused; got " + str(result.get("status"))
        if not _passed(result):
            return False, "matching table refused: " + str(result.get("status"))
        return True, ""
    return Case("XLSX_VERIFY", goal, oracle, not tamper, [])


# ----------------------------------------------------------------------------- bm25 ranking

def gen_source_ranking_bm25(rng: random.Random, work: Path, i: int) -> Case:
    work.mkdir(parents=True, exist_ok=True)
    term1, term2 = _word(rng, 5), _word(rng, 5)
    relevant_doc = {
        "title": f"Guide to {term1} and {term2}",
        "snippet": f"Detailed documentation covering {term1} architecture and {term2} performance tuning."
    }
    irrelevant_doc = {
        "title": "Unrelated topic",
        "snippet": "Nothing to see here about other things entirely different from the query."
    }
    src = work / f"bm25_in_{i}_{_word(rng, 4)}.json"
    out = work / f"bm25_out_{i}_{_word(rng, 4)}.json"
    src.write_text(json.dumps({
        "objective": f"find {term1} {term2} documentation",
        "candidates": [irrelevant_doc, relevant_doc]
    }), encoding="utf-8")
    goal = rng.choice([
        f"Rank candidate sources in {_rel(src)} for objective relevance with BM25 and save to {_rel(out)}.",
        f"Rank the candidate sources in {_rel(src)} for objective relevance with BM25 and save the ranking to {_rel(out)}.",
        f"With BM25, rank candidate sources in {_rel(src)} for objective relevance and save to {_rel(out)}.",
    ])

    def oracle(result: dict[str, Any]) -> tuple[bool, str]:
        if not _passed(result):
            return False, "runtime did not pass: " + str(result.get("status"))
        if not out.is_file():
            return False, "output not written"
        data = json.loads(out.read_text(encoding="utf-8"))
        if data.get("status") != "LEXICAL_RELEVANCE_RANKED":
            return False, f"unexpected status: {data.get('status')}"
        ranked = data.get("ranked_candidates") or []
        if not ranked or ranked[0].get("original_index") != 1:
            return False, f"expected candidate 1 on top, got: {[r.get('original_index') for r in ranked]}"
        return True, ""
    return Case("SOURCE_RANKING_BM25", goal, oracle, True, [out])


# ----------------------------------------------------------------------------- decision synthesis

def gen_decision_synthesis(rng: random.Random, work: Path, i: int) -> Case:
    work.mkdir(parents=True, exist_ok=True)
    alt1, alt2 = _word(rng, 5), _word(rng, 5)
    problem = {
        "schema": "PROJECT_BRAIN_TYPED_DECISION_PROBLEM_V1",
        "decision_id": f"D_{i}_{_word(rng, 3)}",
        "question": f"Which option to select: {alt1} or {alt2}?",
        "alternatives": [
            {"id": alt1, "label": alt1.capitalize()},
            {"id": alt2, "label": alt2.capitalize()},
        ],
        "constraints": [],
        "evidence": [
            {
                "id": "E1",
                "alternative": alt1,
                "relation": "SUPPORTS",
                "confidence": 0.95,
                "strength": 1.0,
                "independence_group": "group_a",
                "claim": f"{alt1} satisfies all requirements",
                "provenance": [{"source": "benchmark_data", "ref": "rec_001"}]
            }
        ],
        "policy": {
            "minimum_support_ratio": 0.60,
            "maximum_conflict_ratio": 0.50,
            "minimum_independent_support_groups": 1,
            "minimum_net_margin": 0.10
        }
    }
    src = work / f"decision_prob_{i}_{_word(rng, 4)}.json"
    out = work / f"decision_res_{i}_{_word(rng, 4)}.json"
    src.write_text(json.dumps(problem), encoding="utf-8")
    goal = rng.choice([
        f"Synthesize a typed decision from {_rel(src)} and save the result to {_rel(out)}.",
        f"Synthesize a typed decision from {_rel(src)} and save it to {_rel(out)}.",
        f"Derive a typed decision from {_rel(src)} and save the result to {_rel(out)}.",
    ])

    def oracle(result: dict[str, Any]) -> tuple[bool, str]:
        if not _passed(result):
            return False, "runtime did not pass: " + str(result.get("status"))
        if not out.is_file():
            return False, "decision result not written"
        data = json.loads(out.read_text(encoding="utf-8"))
        if data.get("selected_alternative") != alt1:
            return False, f"expected {alt1} selected, got {data.get('selected_alternative')}"
        return True, ""
    return Case("DECISION_SYNTHESIS", goal, oracle, True, [out])


# ----------------------------------------------------------------------------- python source audit

def gen_python_source_audit(rng: random.Random, work: Path, i: int) -> Case:
    work.mkdir(parents=True, exist_ok=True)
    src_dir = work / f"src_{i}_{_word(rng, 4)}"
    src_dir.mkdir(parents=True, exist_ok=True)
    fn_name = "calc_" + _word(rng, 5)
    cls_name = "Model_" + _word(rng, 4).capitalize()
    (src_dir / f"mod_{_word(rng, 3)}.py").write_text(
        f"import os\n\ndef {fn_name}(x):\n    return x * 2\n\nclass {cls_name}:\n    pass\n",
        encoding="utf-8"
    )
    out = work / f"audit_{i}_{_word(rng, 4)}.json"
    goal = rng.choice([
        f"Audit the python sources in {_rel(src_dir)} with ast and save the report to {_rel(out)}.",
        f"Audit the python sources in {_rel(src_dir)} with the ast audit and save the report to {_rel(out)}.",
        f"Run the ast audit on the python sources in {_rel(src_dir)} and save the report to {_rel(out)}.",
    ])

    def oracle(result: dict[str, Any]) -> tuple[bool, str]:
        if not _passed(result):
            return False, "runtime did not pass: " + str(result.get("status"))
        if not out.is_file():
            return False, "report not written"
        data = json.loads(out.read_text(encoding="utf-8"))
        totals = data.get("totals") or {}
        if totals.get("files") != 1 or totals.get("functions") != 1:
            return False, f"unexpected totals: {totals}"
        files = data.get("files") or []
        if not files or fn_name not in files[0].get("functions", []):
            return False, f"function {fn_name} not found in audit"
        return True, ""
    return Case("PYTHON_SOURCE_AUDIT", goal, oracle, True, [out])


GENERATORS_V4: dict[str, Callable[[random.Random, Path, int], Case]] = {
    "DOCX_CREATE": gen_docx_create,
    "XLSX_CREATE": gen_xlsx_create,
    "XLSX_VERIFY": gen_xlsx_verify,
    "SOURCE_RANKING_BM25": gen_source_ranking_bm25,
    "DECISION_SYNTHESIS": gen_decision_synthesis,
    "PYTHON_SOURCE_AUDIT": gen_python_source_audit,
}
