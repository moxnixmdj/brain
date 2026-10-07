"""Fail-closed compiler and verifier for a narrow objective subset of GDPval rubric criteria.

Only whole-criterion forms with direct machine-observable artifact semantics are admitted.
Unsupported or mixed semantic criteria remain unresolved. No quality, adequacy, or terminal
credit is implied by compilation alone.
"""
from __future__ import annotations

from dataclasses import dataclass, asdict
from pathlib import Path
import re
import wave
import zipfile
from typing import Any, Mapping

SCHEMA = "PROJECT_BRAIN_GDPVAL_OBJECTIVE_ARTIFACT_CRITERIA_V1"


@dataclass(frozen=True)
class Check:
    kind: str
    args: dict[str, Any]


def _norm_quote(s: str) -> str:
    return s.strip().strip("'\"“”‘’")


def _criterion_payload(criterion: str, criterion_id: str | None = None) -> dict[str, Any]:
    import hashlib
    return {
        "criterion_id": criterion_id,
        "criterion": criterion,
        "criterion_sha256": hashlib.sha256(criterion.encode("utf-8")).hexdigest(),
    }


def compile_criterion(criterion: str, *, criterion_id: str | None = None) -> dict[str, Any]:
    text = " ".join(str(criterion or "").split())
    base = {"schema": SCHEMA, **_criterion_payload(text, criterion_id), "terminal_authority": False}
    if not text:
        return {**base, "status": "FAIL_CLOSED", "reason": "CRITERION_MISSING"}

    checks: list[Check] = []

    # Exact artifact filename. Admission is intentionally limited to native formats
    # whose validity this verifier can independently reopen.
    m = re.fullmatch(
        r"(?:The\s+)?(?:submitted\s+)?(?:workbook|deliverable|audio)?\s*file(?:'s)?\s*"
        r"(?:name|filename)\s+is\s+exactly\s+(.+?)[.]?",
        text,
        re.I,
    )
    if m:
        filename = _norm_quote(m.group(1).rstrip("."))
        if (
            filename
            and Path(filename).name == filename
            and Path(filename).suffix.lower()
            in {".xlsx", ".xlsm", ".xls", ".docx", ".pptx", ".pdf", ".zip", ".wav"}
        ):
            checks.append(Check("EXACT_FILENAME_NATIVE_VALID", {"filename": filename}))

    # Exact named audio delivery. Filename alone is insufficient; the matched WAV
    # must also reopen as a valid WAV artifact.
    if not checks:
        m = re.fullmatch(
            r"An\s+audio\s+file\s+named\s+exactly\s+[\"“](.+?)[\"”]\s+is\s+delivered[.]?",
            text,
            re.I,
        )
        if m and Path(m.group(1)).suffix.lower() == ".wav":
            checks.append(Check("EXACT_FILENAME_NATIVE_VALID", {"filename": m.group(1)}))

    # Workbook basename with allowed Excel extensions.
    if not checks:
        m = re.fullmatch(
            r"The\s+submitted\s+deliverable\s+is\s+an\s+Excel\s+workbook\s+file\s+whose\s+"
            r"basename\s+is\s+['\"“](.+?)['\"”]\s*\(accept\s+\.xlsx,\s*\.xls,\s*or\s*\.xlsm\)[.]?",
            text,
            re.I,
        )
        if m:
            checks.append(
                Check(
                    "EXACT_STEM_ALLOWED_EXTENSIONS",
                    {"stem": m.group(1), "extensions": [".xls", ".xlsm", ".xlsx"]},
                )
            )

    # Only role-free whole criteria are admitted here. Criteria such as
    # "one checklist PDF" remain unresolved because file type does not prove
    # semantic role identity.
    if not checks:
        single_patterns = (
            (
                r"Exactly\s+one\s+top-level\s+ZIP\s+archive\s+is\s+submitted\s+as\s+the\s+"
                r"final\s+deliverable\s*\(one\s+and\s+only\s+one\s+\.zip\s+file\)[.]?",
                ".zip",
            ),
            (
                r"Provides\s+the\s+deliverable\s+as\s+a\s+single\s+Microsoft\s+Excel\s+"
                r"workbook\s+in\s+\.xlsx\s+format[.]?",
                ".xlsx",
            ),
        )
        for pat, ext in single_patterns:
            if re.fullmatch(pat, text, re.I):
                checks.append(
                    Check(
                        "EXACT_TOP_LEVEL_FILE_COUNT_BY_EXTENSION",
                        {
                            "count": 1,
                            "extension": ext,
                            "total_files_must_equal_count": True,
                        },
                    )
                )
                break

    # Simple file-format requirements without extra semantic clauses.
    if not checks:
        format_patterns = (
            (
                r"The\s+final\s+deliverable\s+is\s+provided\s+as\s+an\s+Excel\s+workbook\s+"
                r"in\s+\.xlsx\s+format[.]?",
                ".xlsx",
            ),
            (
                r"Deliverable\s+is\s+provided\s+as\s+a\s+Microsoft\s+Word\s+document\s*"
                r"\(\.docx\)[.]?",
                ".docx",
            ),
        )
        for pat, ext in format_patterns:
            if re.fullmatch(pat, text, re.I):
                checks.append(Check("AT_LEAST_ONE_VALID_FILE_BY_EXTENSION", {"extension": ext}))
                break

    # Additional exact whole-criterion mechanical forms observed in the frozen
    # public-220 rubric. These are intentionally lexical and conservative:
    # semantic role labels or mixed content requirements are not admitted.
    if not checks:
        exact_single_ext_patterns = (
            (
                r"Delivers\s+a\s+single\s+Excel\s+workbook\s+file\s+in\s+\.xlsx\s+format[.]?",
                ".xlsx",
            ),
            (
                r"Output\s+is\s+a\s+single\s+Excel\s+workbook\s+file\s+with\s+the\s+\.xlsx\s+extension[.]?",
                ".xlsx",
            ),
            (
                r"The\s+submitted\s+file\s+is\s+an\s+\.xlsx\s+Excel\s+workbook\s+"
                r"\(not\s+\.xlsm\s+or\s+other\s+formats\)[.]?",
                ".xlsx",
            ),
            (
                r"The\s+deliverable\s+is\s+a\s+single\s+Word\s+document\s*\(\.docx\)\s+"
                r"rather\s+than\s+multiple\s+files\s+or\s+a\s+PDF[.]?",
                ".docx",
            ),
            (
                r"Exactly\s+one\s+Microsoft\s+Word\s+file\s+is\s+delivered\s+and\s+it\s+opens\s+"
                r"as\s+a\s+valid\s+DOCX\s+package\s+\(ZIP\s+with\s+\[Content_Types\]\.xml\s+"
                r"and\s+word/document\.xml\)[.]?",
                ".docx",
            ),
            (
                r"Submission\s+consists\s+of\s+exactly\s+one\s+file\s+in\s+PDF\s+format\s*"
                r"\(\.pdf\)[.]?",
                ".pdf",
            ),
            (
                r"(?:The\s+)?(?:final\s+)?(?:deliverable|submission)\s+(?:is|is\s+provided\s+as|"
                r"is\s+provided\s+as\s+a|is\s+provided\s+as\s+a\s+single|"
                r"is\s+a|is\s+a\s+single|is\s+provided\s+as\s+single)\s+"
                r"(?:single\s+)?PDF\s+(?:file|document)[.]?",
                ".pdf",
            ),
        )
        for pat, ext in exact_single_ext_patterns:
            if re.fullmatch(pat, text, re.I):
                checks.append(
                    Check(
                        "EXACT_TOP_LEVEL_FILE_COUNT_BY_EXTENSION",
                        {
                            "count": 1,
                            "extension": ext,
                            "total_files_must_equal_count": True,
                        },
                    )
                )
                break

    # Simple role-free artifact-format statements. Requiring a natively
    # reopenable matching file is a sound one-sided check; extra files do not
    # falsify these criteria unless the criterion explicitly says "single".
    if not checks:
        simple_ext_patterns = (
            (r"(?:The\s+)?(?:final\s+)?deliverable\s+is\s+(?:provided\s+as\s+)?a\s+PDF(?:\s+file|\s+document)?[.]?", ".pdf"),
            (r"The\s+deliverable\s+is\s+provided\s+as\s+a\s+Microsoft\s+Word\s+document\s+with\s+a\s+\.docx\s+file\s+extension[.]?", ".docx"),
            (r"The\s+deliverable\s+is\s+provided\s+as\s+a\s+\.pptx\s+PowerPoint\s+file[.]?", ".pptx"),
        )
        for pat, ext in simple_ext_patterns:
            if re.fullmatch(pat, text, re.I):
                checks.append(Check("AT_LEAST_ONE_VALID_FILE_BY_EXTENSION", {"extension": ext}))
                break

    # Exact single-page PDF submission.
    if not checks and re.fullmatch(
        r"(?:The\s+)?Submission\s+is\s+a\s+single-page\s+PDF[.]?",
        text,
        re.I,
    ):
        checks.append(
            Check(
                "SINGLE_PDF_PAGE_COUNT",
                {"page_count": 1, "total_files_must_equal_count": True},
            )
        )

    # Exact workbook sheet count, with no semantic sheet-role inference.
    if not checks:
        m = re.fullmatch(
            r"Provides\s+an\s+Excel\s+\.xlsx\s+workbook\s+with\s+exactly\s+(\d+)\s+"
            r"worksheets\s*\(tabs\)[.]?",
            text,
            re.I,
        )
        if m:
            checks.append(Check("SINGLE_XLSX_SHEET_COUNT", {"sheet_count": int(m.group(1))}))

    # Exact uncompressed WAV fact. We deliberately require one top-level file,
    # which is stricter than merely accepting an arbitrary matching WAV among
    # unrelated files and therefore cannot create a false positive.
    if not checks and re.fullmatch(
        r"The\s+delivered\s+file\s+is\s+a\s+WAV\s+file\s+and\s+not\s+a\s+compressed\s+"
        r"format\s*\(e\.g\.,\s*MP3/AAC\)[.]?",
        text,
        re.I,
    ):
        checks.append(Check("SINGLE_WAV_UNCOMPRESSED", {"total_files_must_equal_count": True}))

    # Single Word or PDF document.
    if not checks and re.fullmatch(
        r"The\s+deliverable\s+is\s+provided\s+as\s+a\s+single\s+Word\s*\(\.docx\)\s+or\s+"
        r"PDF\s*\(\.pdf\)\s+document[.]?",
        text,
        re.I,
    ):
        checks.append(
            Check(
                "EXACT_TOP_LEVEL_FILE_COUNT_ALLOWED_EXTENSIONS",
                {
                    "count": 1,
                    "extensions": [".docx", ".pdf"],
                    "total_files_must_equal_count": True,
                },
            )
        )

    # Worksheet existence/name with exact case-fold + surrounding-space semantics.
    if not checks:
        m = re.fullmatch(
            r"The\s+workbook\s+contains\s+a\s+worksheet\s+named\s+exactly\s+['\"“](.+?)['\"”]\s*"
            r"\(case-insensitive,\s*ignoring\s+surrounding\s+spaces\)[.]?",
            text,
            re.I,
        )
        if m:
            checks.append(Check("XLSX_SHEET_NAME_CASEFOLD_TRIM", {"sheet_name": m.group(1)}))

    # Native reopen validity, only when the whole criterion is exactly that fact.
    if not checks:
        m = re.fullmatch(r"The\s+\.(docx|xlsx|pptx)\s+opens\s+without\s+error[.]?", text, re.I)
        if m:
            checks.append(
                Check("AT_LEAST_ONE_VALID_FILE_BY_EXTENSION", {"extension": "." + m.group(1).lower()})
            )

    # WAV header facts. Conditional variants are deliberately unresolved.
    if not checks:
        m = re.fullmatch(
            r"The\s+delivered\s+audio\s+file\s+has\s+a\s+sampling\s+rate\s+of\s+exactly\s+"
            r"48,000\s+Hz[.]?",
            text,
            re.I,
        )
        if m:
            checks.append(Check("SINGLE_WAV_SAMPLE_RATE", {"sample_rate_hz": 48000}))

        m2 = re.fullmatch(
            r"The\s+delivered\s+audio\s+file\s+has\s+a\s+bit\s+depth\s+of\s+exactly\s+24\s+"
            r"bits\s+per\s+sample\s*\(PCM-equivalent\s+on\s+decode\)[.]?",
            text,
            re.I,
        )
        if m2:
            checks.append(Check("SINGLE_WAV_SAMPLE_WIDTH_BITS", {"bits": 24}))

    if not checks:
        return {
            **base,
            "status": "UNRESOLVED",
            "reason": "OUTSIDE_EXACT_OBJECTIVE_ARTIFACT_GRAMMAR",
        }

    return {
        **base,
        "status": "COMPILED_OBJECTIVE_CHECK",
        "checks": [asdict(c) for c in checks],
        "soundness_boundary": (
            "WHOLE_CRITERION_EXACT_MATCH_ONLY__NO_IGNORED_CONJUNCTS__"
            "ROLE_IDENTITY_NOT_INFERRED_FROM_FILE_TYPE__"
            "NO_QUALITY_OR_SEMANTIC_EQUIVALENCE_INFERENCE"
        ),
    }


def _files(root: Path) -> list[Path]:
    return sorted([p for p in root.iterdir() if p.is_file()], key=lambda p: p.name.casefold())


def _valid_package(path: Path) -> bool:
    """Reopen through the same native libraries used by Brain-owned artifact mechanics."""
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            from docx import Document

            Document(str(path))
            return True

        if ext in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            wb = load_workbook(
                str(path),
                read_only=True,
                data_only=False,
                keep_vba=(ext == ".xlsm"),
            )
            try:
                return bool(wb.sheetnames)
            finally:
                wb.close()

        if ext == ".pptx":
            from pptx import Presentation

            Presentation(str(path))
            return True

        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path), strict=True)
            _ = len(reader.pages)
            return True

        if ext == ".zip":
            with zipfile.ZipFile(path) as z:
                return z.testzip() is None

        if ext == ".wav":
            with wave.open(str(path), "rb") as w:
                return w.getnchannels() > 0 and w.getframerate() > 0 and w.getsampwidth() > 0

        if ext == ".xls":
            # Legacy binary XLS remains a deliberate false negative until a
            # Brain-owned native XLS parser is separately bound.
            return False
    except Exception:
        return False

    return path.exists() and path.stat().st_size > 0


def _xlsx_sheet_names(path: Path) -> list[str] | None:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(
            str(path),
            read_only=True,
            data_only=False,
            keep_vba=(path.suffix.lower() == ".xlsm"),
        )
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
    except Exception:
        return None


def verify_compiled(compiled: Mapping[str, Any], artifact_dir: str | Path) -> dict[str, Any]:
    if compiled.get("status") != "COMPILED_OBJECTIVE_CHECK":
        return {
            "schema": SCHEMA,
            "status": "FAIL_CLOSED",
            "pass": False,
            "errors": ["CRITERION_NOT_COMPILED"],
            "terminal_authority": False,
        }

    root = Path(artifact_dir)
    if not root.is_dir():
        return {
            "schema": SCHEMA,
            "status": "FAIL_CLOSED",
            "pass": False,
            "errors": ["ARTIFACT_DIR_INVALID"],
            "terminal_authority": False,
        }

    files = _files(root)
    errors: list[str] = []
    evidence: list[dict[str, Any]] = []

    for row in compiled.get("checks", []):
        kind = row.get("kind")
        args = row.get("args") or {}

        if kind == "EXACT_FILENAME_NATIVE_VALID":
            cand = [p for p in files if p.name == args.get("filename")]
            ok = len(cand) == 1 and _valid_package(cand[0])

        elif kind == "EXACT_STEM_ALLOWED_EXTENSIONS":
            exts = set(args.get("extensions") or [])
            ok = any(
                p.stem == args.get("stem")
                and p.suffix.lower() in exts
                and _valid_package(p)
                for p in files
            )

        elif kind == "EXACT_TOP_LEVEL_FILE_COUNT_BY_EXTENSION":
            ext = str(args.get("extension") or "").lower()
            cand = [p for p in files if p.suffix.lower() == ext]
            target_count = int(args.get("count", -1))
            ok = (
                len(cand) == target_count
                and all(_valid_package(p) for p in cand)
                and (
                    not args.get("total_files_must_equal_count")
                    or len(files) == target_count
                )
            )

        elif kind == "EXACT_TOP_LEVEL_FILE_COUNT_ALLOWED_EXTENSIONS":
            exts = set(args.get("extensions") or [])
            cand = [p for p in files if p.suffix.lower() in exts]
            target_count = int(args.get("count", -1))
            ok = (
                len(cand) == target_count
                and all(_valid_package(p) for p in cand)
                and (
                    not args.get("total_files_must_equal_count")
                    or len(files) == target_count
                )
            )

        elif kind == "AT_LEAST_ONE_VALID_FILE_BY_EXTENSION":
            ext = str(args.get("extension") or "").lower()
            cand = [p for p in files if p.suffix.lower() == ext]
            ok = bool(cand) and all(_valid_package(p) for p in cand)

        elif kind == "XLSX_SHEET_NAME_CASEFOLD_TRIM":
            target = str(args.get("sheet_name") or "").strip().casefold()
            ok = False
            for p in files:
                if p.suffix.lower() not in {".xlsx", ".xlsm"}:
                    continue
                names = _xlsx_sheet_names(p)
                if names is not None and any(n.strip().casefold() == target for n in names):
                    ok = True
                    break

        elif kind == "SINGLE_PDF_PAGE_COUNT":
            cand = [p for p in files if p.suffix.lower() == ".pdf"]
            target_count = int(args.get("page_count", -1))
            ok = False
            if len(cand) == 1 and (
                not args.get("total_files_must_equal_count") or len(files) == 1
            ):
                try:
                    from pypdf import PdfReader

                    reader = PdfReader(str(cand[0]), strict=True)
                    ok = len(reader.pages) == target_count
                except Exception:
                    ok = False

        elif kind == "SINGLE_XLSX_SHEET_COUNT":
            cand = [p for p in files if p.suffix.lower() == ".xlsx"]
            target_count = int(args.get("sheet_count", -1))
            if len(cand) != 1:
                ok = False
            else:
                names = _xlsx_sheet_names(cand[0])
                ok = names is not None and len(names) == target_count

        elif kind == "SINGLE_WAV_UNCOMPRESSED":
            cand = [p for p in files if p.suffix.lower() == ".wav"]
            ok = False
            if len(cand) == 1 and (
                not args.get("total_files_must_equal_count") or len(files) == 1
            ):
                try:
                    with wave.open(str(cand[0]), "rb") as w:
                        ok = w.getcomptype() == "NONE"
                except Exception:
                    ok = False

        elif kind == "SINGLE_WAV_SAMPLE_RATE":
            cand = [p for p in files if p.suffix.lower() == ".wav"]
            if len(cand) != 1:
                ok = False
            else:
                try:
                    with wave.open(str(cand[0]), "rb") as w:
                        ok = w.getframerate() == int(args.get("sample_rate_hz"))
                except Exception:
                    ok = False

        elif kind == "SINGLE_WAV_SAMPLE_WIDTH_BITS":
            cand = [p for p in files if p.suffix.lower() == ".wav"]
            if len(cand) != 1:
                ok = False
            else:
                try:
                    with wave.open(str(cand[0]), "rb") as w:
                        ok = w.getsampwidth() * 8 == int(args.get("bits"))
                except Exception:
                    ok = False

        else:
            errors.append("UNKNOWN_CHECK_KIND:" + str(kind))
            continue

        evidence.append({"kind": kind, "pass": bool(ok)})
        if not ok:
            errors.append("CHECK_FAILED:" + str(kind))

    return {
        "schema": SCHEMA,
        "status": "PASS__OBJECTIVE_CRITERION_VERIFIED" if not errors else "FALSIFIED_OR_FAIL_CLOSED",
        "pass": not errors,
        "criterion_id": compiled.get("criterion_id"),
        "criterion_sha256": compiled.get("criterion_sha256"),
        "evidence": evidence,
        "errors": errors,
        "terminal_authority": False,
        "u_deletion_authority": False,
    }
