import tempfile
import wave
import zipfile
from pathlib import Path

from canonical.runtime import gdpval_objective_artifact_criteria_v1 as m


def test_exact_workbook_filename_requires_native_validity():
    c = m.compile_criterion(
        "The submitted workbook file name is exactly Aurisic_Financials_4-25-1.xlsx."
    )
    assert c["status"] == "COMPILED_OBJECTIVE_CHECK"
    with tempfile.TemporaryDirectory() as d:
        from openpyxl import Workbook

        p = Path(d) / "Aurisic_Financials_4-25-1.xlsx"
        wb = Workbook()
        wb.save(p)
        wb.close()
        assert m.verify_compiled(c, d)["pass"] is True

        p.write_bytes(b"not-an-xlsx")
        assert m.verify_compiled(c, d)["pass"] is False


def test_sheet_name_casefold_trim():
    c = m.compile_criterion(
        "The workbook contains a worksheet named exactly 'Sample Size Calculation' "
        "(case-insensitive, ignoring surrounding spaces)."
    )
    assert c["status"] == "COMPILED_OBJECTIVE_CHECK"
    with tempfile.TemporaryDirectory() as d:
        from openpyxl import Workbook

        p = Path(d) / "a.xlsx"
        wb = Workbook()
        wb.active.title = "  sample size calculation "
        wb.save(p)
        wb.close()
        assert m.verify_compiled(c, d)["pass"] is True


def test_mixed_semantics_refused():
    assert m.compile_criterion(
        "Both .pptx files open without error and are not password-protected."
    )["status"] == "UNRESOLVED"
    assert m.compile_criterion(
        "The Excel schedule preserves the rotation order from the source PDF."
    )["status"] == "UNRESOLVED"


def test_semantic_role_labeled_file_count_refused():
    assert m.compile_criterion(
        "Exactly one checklist file is provided in PDF format."
    )["status"] == "UNRESOLVED"
    assert m.compile_criterion(
        "Exactly one action-items document is provided in Word (.docx) format."
    )["status"] == "UNRESOLVED"


def test_wav_header():
    c = m.compile_criterion(
        "The delivered audio file has a sampling rate of exactly 48,000 Hz."
    )
    assert c["status"] == "COMPILED_OBJECTIVE_CHECK"
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "a.wav"
        with wave.open(str(p), "wb") as w:
            w.setnchannels(1)
            w.setsampwidth(3)
            w.setframerate(48000)
            w.writeframes(b"\x00\x00\x00" * 10)
        assert m.verify_compiled(c, d)["pass"] is True


def test_exact_audio_filename_requires_valid_wav():
    c = m.compile_criterion(
        'An audio file named exactly "State of Affairs_FULL_EDIT_MIX.wav" is delivered.'
    )
    assert c["status"] == "COMPILED_OBJECTIVE_CHECK"
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "State of Affairs_FULL_EDIT_MIX.wav"
        p.write_bytes(b"fake")
        assert m.verify_compiled(c, d)["pass"] is False

        with wave.open(str(p), "wb") as w:
            w.setnchannels(1)
            w.setsampwidth(2)
            w.setframerate(48000)
            w.writeframes(b"\x00\x00" * 4)
        assert m.verify_compiled(c, d)["pass"] is True


def test_simple_docx_format_and_corruption():
    c = m.compile_criterion(
        "Deliverable is provided as a Microsoft Word document (.docx)."
    )
    assert c["status"] == "COMPILED_OBJECTIVE_CHECK"
    with tempfile.TemporaryDirectory() as d:
        from docx import Document

        p = Path(d) / "x.docx"
        doc = Document()
        doc.add_paragraph("x")
        doc.save(p)
        assert m.verify_compiled(c, d)["pass"] is True

        p.write_bytes(b"not-a-docx")
        assert m.verify_compiled(c, d)["pass"] is False


def test_professional_quality_refused():
    assert m.compile_criterion(
        "The memo is clear and professional."
    )["status"] == "UNRESOLVED"


def test_single_zip_rejects_extra_top_level_file():
    c = m.compile_criterion(
        "Exactly one top-level ZIP archive is submitted as the final deliverable "
        "(one and only one .zip file)."
    )
    assert c["status"] == "COMPILED_OBJECTIVE_CHECK"
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "bundle.zip"
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("x.txt", "x")
        (Path(d) / "extra.txt").write_text("x")
        assert m.verify_compiled(c, d)["pass"] is False


def test_frozen_public_220_simple_pdf_forms_compile_without_role_inference():
    for criterion in [
        "The deliverable is a single PDF file.",
        "The deliverable is provided as a single PDF document",
        "Submission consists of exactly one file in PDF format (.pdf).",
        "The deliverable is provided as a PDF file.",
        "Deliverable is a PDF.",
    ]:
        out = m.compile_criterion(criterion)
        assert out["status"] == "COMPILED_OBJECTIVE_CHECK", criterion


def test_single_page_pdf_is_machine_verified(tmp_path):
    from pypdf import PdfWriter

    c = m.compile_criterion("Submission is a single-page PDF")
    assert c["status"] == "COMPILED_OBJECTIVE_CHECK"

    p = tmp_path / "one.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with p.open("wb") as f:
        writer.write(f)
    assert m.verify_compiled(c, tmp_path)["pass"] is True

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_blank_page(width=100, height=100)
    with p.open("wb") as f:
        writer.write(f)
    assert m.verify_compiled(c, tmp_path)["pass"] is False


def test_exact_four_sheet_xlsx_is_machine_verified(tmp_path):
    from openpyxl import Workbook

    c = m.compile_criterion(
        "Provides an Excel .xlsx workbook with exactly 4 worksheets (tabs)."
    )
    assert c["status"] == "COMPILED_OBJECTIVE_CHECK"

    p = tmp_path / "book.xlsx"
    wb = Workbook()
    for name in ["B", "C", "D"]:
        wb.create_sheet(name)
    wb.save(p)
    wb.close()
    assert m.verify_compiled(c, tmp_path)["pass"] is True


def test_new_mechanical_forms_stay_fail_closed_on_mixed_semantics():
    assert m.compile_criterion(
        "The deliverable is a single PDF file containing a complete legal analysis."
    )["status"] == "UNRESOLVED"
    assert m.compile_criterion(
        "Provides two distinct .pptx files: one presentation for Session 13 and one for Session 14."
    )["status"] == "UNRESOLVED"


def test_frozen_public_220_exact_single_native_forms_compile():
    for criterion in [
        "Delivers a single Excel workbook file in .xlsx format.",
        "Output is a single Excel workbook file with the .xlsx extension.",
        "The submitted file is an .xlsx Excel workbook (not .xlsm or other formats).",
        "The deliverable is a single Word document (.docx) rather than multiple files or a PDF.",
        "Exactly one Microsoft Word file is delivered and it opens as a valid DOCX package (ZIP with [Content_Types].xml and word/document.xml).",
        "The delivered file is a WAV file and not a compressed format (e.g., MP3/AAC).",
    ]:
        out = m.compile_criterion(criterion)
        assert out["status"] == "COMPILED_OBJECTIVE_CHECK", criterion
