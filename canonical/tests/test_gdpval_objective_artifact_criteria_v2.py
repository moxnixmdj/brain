from pathlib import Path

from canonical.runtime import gdpval_objective_artifact_criteria_v1 as m


def test_exact_one_deliverable_file_compiles_to_count_only(tmp_path):
    c = m.compile_criterion("Exactly one deliverable file is submitted.")
    assert c["status"] == "COMPILED_OBJECTIVE_CHECK"
    assert c["checks"] == [{"kind": "EXACT_TOP_LEVEL_FILE_COUNT", "args": {"count": 1}}]

    (tmp_path / "x.txt").write_text("x", encoding="utf-8")
    assert m.verify_compiled(c, tmp_path)["pass"] is True
    (tmp_path / "y.txt").write_text("y", encoding="utf-8")
    assert m.verify_compiled(c, tmp_path)["pass"] is False


def test_simple_docx_extension_form_requires_valid_docx(tmp_path):
    from docx import Document

    criterion = "The deliverable is a Microsoft Word document with a .docx file extension"
    c = m.compile_criterion(criterion)
    assert c["status"] == "COMPILED_OBJECTIVE_CHECK"

    p = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("verified")
    doc.save(p)
    assert m.verify_compiled(c, tmp_path)["pass"] is True

    p.write_bytes(b"not-docx")
    assert m.verify_compiled(c, tmp_path)["pass"] is False


def test_simple_excel_file_accepts_only_native_valid_modern_excel(tmp_path):
    from openpyxl import Workbook

    c = m.compile_criterion("The submitted deliverable is an Excel file.")
    assert c["status"] == "COMPILED_OBJECTIVE_CHECK"
    assert c["checks"][0]["kind"] == "AT_LEAST_ONE_VALID_FILE_ALLOWED_EXTENSIONS"

    p = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.save(p)
    wb.close()
    assert m.verify_compiled(c, tmp_path)["pass"] is True

    p.unlink()
    (tmp_path / "legacy.xls").write_bytes(b"legacy")
    assert m.verify_compiled(c, tmp_path)["pass"] is False


def test_role_or_content_conjunctions_remain_unresolved():
    assert m.compile_criterion(
        "Provides two distinct .pptx files: one presentation for Session 13 and one for Session 14."
    )["status"] == "UNRESOLVED"
    assert m.compile_criterion(
        "A single Excel workbook in .xlsx format is provided and contains the required questions."
    )["status"] == "UNRESOLVED"
    assert m.compile_criterion(
        "Provides a Word (.docx) document that contains the Proposed Architecture Summary"
    )["status"] == "UNRESOLVED"


def test_plain_deliverables_without_explicit_files_are_not_counted():
    assert m.compile_criterion("Response provides two deliverables")["status"] == "UNRESOLVED"
