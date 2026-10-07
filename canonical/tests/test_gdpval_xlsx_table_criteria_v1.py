import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from canonical.runtime import gdpval_xlsx_table_criteria_v1 as m


def _book(path:Path, first_headers, first_rows, second_headers=None, second_rows=None):
    wb=Workbook()
    ws=wb.active
    ws.title="First"
    ws.append(first_headers)
    for r in first_rows: ws.append(r)
    ref=f"A1:{chr(64+len(first_headers))}{1+len(first_rows)}"
    tab=Table(displayName="Table1",ref=ref)
    tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True)
    ws.add_table(tab)
    if second_headers is not None:
        ws2=wb.create_sheet("Second")
        ws2.append(second_headers)
        for r in second_rows or []: ws2.append(r)
        ref2=f"A1:{chr(64+len(second_headers))}{1+len(second_rows or [])}"
        tab2=Table(displayName="Table2",ref=ref2)
        tab2.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True)
        ws2.add_table(tab2)
    wb.save(path)
    wb.close()


def test_first_worksheet_native_table_column_casefold_passes():
    c=m.compile_criterion("First worksheet includes a column titled 'Water system' (case-insensitive match)")
    assert c["status"]=="COMPILED_SUFFICIENT_XLSX_TABLE_CHECK"
    with tempfile.TemporaryDirectory() as d:
        _book(Path(d)/"x.xlsx",["WATER SYSTEM","Well ID"],[["A",1]])
        assert m.verify_compiled(c,d)["pass"] is True


def test_second_worksheet_column_requires_second_sheet():
    c=m.compile_criterion("Second worksheet includes a column titled 'Status' (case-insensitive match)")
    with tempfile.TemporaryDirectory() as d:
        _book(Path(d)/"x.xlsx",["Status"],[["ok"]])
        out=m.verify_compiled(c,d)
        assert out["pass"] is False
        assert "WORKSHEET_ORDINAL_MISSING" in out["errors"]


def test_matching_cell_text_without_native_table_is_not_accepted():
    c=m.compile_criterion("First worksheet includes a column titled 'Water system' (case-insensitive match)")
    with tempfile.TemporaryDirectory() as d:
        wb=Workbook()
        ws=wb.active
        ws["A1"]="Water system"
        ws["A2"]="A"
        wb.save(Path(d)/"x.xlsx")
        wb.close()
        out=m.verify_compiled(c,d)
        assert out["pass"] is False
        assert "NATIVE_TABLE_COLUMN_WITNESS_MISSING" in out["errors"]


def test_second_worksheet_global_not_contains_passes():
    c=m.compile_criterion(
        "No row on the second worksheet has Well Description containing 'abandoned' (case-insensitive)"
    )
    with tempfile.TemporaryDirectory() as d:
        _book(
            Path(d)/"x.xlsx",
            ["A"],[[1]],
            ["Well Description","Status"],
            [["active well","ok"],["production","ok"]],
        )
        assert m.verify_compiled(c,d)["pass"] is True


def test_second_worksheet_global_not_contains_catches_value_below_table():
    c=m.compile_criterion(
        "No row on the second worksheet has Well Description containing 'abandoned' (case-insensitive)"
    )
    with tempfile.TemporaryDirectory() as d:
        _book(
            Path(d)/"x.xlsx",
            ["A"],[[1]],
            ["Well Description","Status"],
            [["active well","ok"]],
        )
        from openpyxl import load_workbook
        p=Path(d)/"x.xlsx"
        wb=load_workbook(p)
        ws=wb.worksheets[1]
        ws["A4"]="ABANDONED WELL"
        wb.save(p)
        wb.close()
        out=m.verify_compiled(c,d)
        assert out["pass"] is False
        assert out["detail"]["violations"]


def test_arbitrary_row_title_template_remains_unresolved():
    # Without an explicit worksheet/table identity, this V1 does not guess which
    # table or artifact owns the phrase.
    assert m.compile_criterion(
        'Includes a row titled "Equities" under "Asset Class" column'
    )["status"]=="UNRESOLVED"


def test_reasonable_variant_column_semantics_remain_unresolved():
    assert m.compile_criterion(
        "The Excel table contains a column labeled (case-insensitive, reasonable variants accepted) 'Hardware Type'."
    )["status"]=="UNRESOLVED"
