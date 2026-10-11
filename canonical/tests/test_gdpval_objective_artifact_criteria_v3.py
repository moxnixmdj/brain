from pathlib import Path
import tempfile, wave
from canonical.runtime import gdpval_objective_artifact_criteria_v3 as m
from canonical.runtime.gdpval_first_residual_mechanical_replay_v3 import replay

def _pdf(path:Path,n:int):
 from pypdf import PdfWriter
 w=PdfWriter()
 for _ in range(n):w.add_blank_page(width=100,height=100)
 with path.open("wb") as f:w.write(f)

def _pptx(path:Path,n:int):
 from pptx import Presentation
 p=Presentation()
 # default has zero slides
 for _ in range(n):p.slides.add_slide(p.slide_layouts[6])
 p.save(path)

def _docx(path:Path):
 from docx import Document
 d=Document();d.add_paragraph("x");d.save(path)

def test_frozen_220_v3_exact_gain_and_zero_credit():
 o=replay()
 assert o["v1_count"]==28
 assert o["v2_count"]==92
 assert o["v3_count"]==131
 assert o["v3_new_count"]==39
 assert o["v3_unresolved_count"]==89
 assert o["acceptance_credit_delta"]==0 and o["terminal_credit_delta"]==0
 assert o["terminal_authority"] is False

def test_role_bearing_or_mixed_criteria_remain_unresolved():
 rejected=[
  "Provides an Excel workbook file (.xlsx or .xlsm) as the primary analytical deliverable",
  "The updated quotation is delivered as a single Excel workbook saved with the exact filename Q9749821-revised_including_transport.xlsx.",
  "Exactly one checklist file is provided in PDF format.",
  "A memo file in PDF format ('.pdf') is submitted.",
  "Provides the scenario analysis as a single Excel workbook with .xlsx extension containing all required elements (three scenarios, computations, visual, and summary).",
  "Submits exactly one PDF file containing the legal memorandum",
  "The final submission is a single PDF file containing the screenplay",
  "Provides an updated Excel workbook in .xlsx format containing the vendor table assignments",
 ]
 for x in rejected:
  assert m.compile_criterion(x)["status"]=="UNRESOLVED",x

def test_page_range_is_sound_pdf_subset(tmp_path:Path):
 c=m.compile_criterion("Document length is 2–3 pages.")
 assert c["criterion_to_check_relation"]=="SOUND_SUFFICIENT_SUBSET__FALSE_NEGATIVES_ALLOWED"
 p=tmp_path/"d.pdf";_pdf(p,2);assert m.verify_compiled(c,tmp_path)["pass"] is True
 _pdf(p,4);assert m.verify_compiled(c,tmp_path)["pass"] is False
 (tmp_path/"extra.txt").write_text("x");_pdf(p,3);assert m.verify_compiled(c,tmp_path)["pass"] is False

def test_second_page_range_wording_compiles():
 c=m.compile_criterion("The document length is between 2 and 3 pages (inclusive).")
 assert c["status"]=="COMPILED_OBJECTIVE_CHECK"
 assert c["checks"][0]["kind"]=="SINGLE_PDF_PAGE_COUNT_RANGE"

def test_slide_ceiling_is_sound_pptx_subset(tmp_path:Path):
 c=m.compile_criterion("Submission length is no more than 10 slides total.")
 assert c["criterion_to_check_relation"]=="SOUND_SUFFICIENT_SUBSET__FALSE_NEGATIVES_ALLOWED"
 p=tmp_path/"d.pptx";_pptx(p,10);assert m.verify_compiled(c,tmp_path)["pass"] is True
 _pptx(p,11);assert m.verify_compiled(c,tmp_path)["pass"] is False

def test_two_deliverables_exact_total_count(tmp_path:Path):
 c=m.compile_criterion("Response provides two deliverables")
 (tmp_path/"a.bin").write_bytes(b"a");(tmp_path/"b.bin").write_bytes(b"b")
 assert m.verify_compiled(c,tmp_path)["pass"] is True
 (tmp_path/"c.bin").write_bytes(b"c")
 assert m.verify_compiled(c,tmp_path)["pass"] is False

def test_exact_three_pdfs_requires_three_valid_and_no_extra(tmp_path:Path):
 c=m.compile_criterion("Exactly three separate PDF files are delivered")
 for i in range(3):_pdf(tmp_path/f"{i}.pdf",1)
 assert m.verify_compiled(c,tmp_path)["pass"] is True
 (tmp_path/"x.txt").write_text("x")
 assert m.verify_compiled(c,tmp_path)["pass"] is False

def test_doc_or_docx_subset_requires_valid_single_docx(tmp_path:Path):
 c=m.compile_criterion("The uploaded deliverable is a Microsoft Word document (.docx or .doc).")
 assert c["criterion_to_check_relation"]=="SOUND_SUFFICIENT_SUBSET__FALSE_NEGATIVES_ALLOWED"
 _docx(tmp_path/"x.docx");assert m.verify_compiled(c,tmp_path)["pass"] is True
 (tmp_path/"extra.txt").write_text("x");assert m.verify_compiled(c,tmp_path)["pass"] is False

def test_named_telehealth_workflow_subset_is_name_and_native_validity(tmp_path:Path):
 c=m.compile_criterion('The file "Telehealth Workflow" is delivered in Word format (.doc or .docx).')
 _docx(tmp_path/"Telehealth Workflow.docx");assert m.verify_compiled(c,tmp_path)["pass"] is True
 (tmp_path/"Telehealth Workflow.docx").rename(tmp_path/"Other.docx")
 assert m.verify_compiled(c,tmp_path)["pass"] is False

def test_generic_audio_is_wav_sufficient_subset(tmp_path:Path):
 c=m.compile_criterion("The deliverable is an audio file.")
 assert c["criterion_to_check_relation"]=="SOUND_SUFFICIENT_SUBSET__FALSE_NEGATIVES_ALLOWED"
 p=tmp_path/"a.wav"
 with wave.open(str(p),"wb") as w:
  w.setnchannels(1);w.setsampwidth(2);w.setframerate(16000);w.writeframes(b"\x00\x00"*4)
 assert m.verify_compiled(c,tmp_path)["pass"] is True
